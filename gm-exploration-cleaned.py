"""Part 2 analysis for the Global Markets case study.

Run the setup block once, then use "Run Below" to run Queries 1 to 8.
Each query writes its output to GM-Query-Outputs-Cleaned.xlsx immediately.
"""

# %% SETUP: libraries, source data and Excel output functions
from io import BytesIO
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from matplotlib.ticker import FuncFormatter
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font

from plot_theme import CHART_COLORS, apply_chart_theme

apply_chart_theme()

BASE_DIR = Path(__file__).resolve().parent
SOURCE_XLSX = BASE_DIR / 'GM-Case-study-Data-2026.xlsx'
RATES_CSV = BASE_DIR / 'Source-Material/exchange-rates-zar-2020-2022.csv'
TARGET_XLSX = BASE_DIR / 'GM-Presentation/GM-Query-Outputs-Cleaned.xlsx'

YEARS = [2020, 2021, 2022]
MONTHS = [
    'January', 'February', 'March', 'April', 'May', 'June',
    'July', 'August', 'September', 'October', 'November', 'December',
]
REVENUE = 'Client Value ZAR'
VOLUME = 'Volume USD'
CUSTOMER = 'Customer Name'
SECTOR = 'Immediate Parent Client Sector'


def join_unique(values):
    return ', '.join(dict.fromkeys(values.dropna()))


def client_dimensions(source):
    return source.groupby(CUSTOMER).agg(
        Sector=(SECTOR, join_unique),
        **{
            'Trading Business Area': ('Trading Business Area', join_unique),
            'Parent Product': ('Parent Product', join_unique),
        },
    )


def margin_bps(revenue, volume):
    return revenue.div(volume.where(volume > 0)).mul(10000)


def check_close(actual, expected, label):
    if not np.allclose(actual, expected, rtol=1e-10, atol=0.01):
        raise ValueError(f'Reconciliation failed: {label}')


def print_table(title, table):
    print(f'\n{title}')
    print(table.to_string(index=False, float_format=lambda x: f'{x:,.2f}', na_rep='N/A'))


def number_rows(table):
    table = table.copy()
    table.insert(0, 'Rank', range(1, len(table) + 1))
    return table


def send_to_xlsx(sheets):
    """Write or replace one query's sheets immediately."""
    TARGET_XLSX.parent.mkdir(parents=True, exist_ok=True)
    file_exists = TARGET_XLSX.exists()
    writer_args = {
        'path': TARGET_XLSX,
        'engine': 'openpyxl',
        'mode': 'a' if file_exists else 'w',
    }
    if file_exists:
        writer_args['if_sheet_exists'] = 'replace'

    with pd.ExcelWriter(**writer_args) as writer:
        for sheet_name, table in sheets.items():
            if len(sheet_name) > 31:
                raise ValueError(f'Excel sheet name is too long: {sheet_name}')
            table.to_excel(writer, sheet_name=sheet_name, index=False)
            ws = writer.sheets[sheet_name]
            ws.freeze_panes = 'A2'
            ws.auto_filter.ref = ws.dimensions
            for column in ws.columns:
                width = max(len(str(cell.value or '')) for cell in column) + 2
                ws.column_dimensions[column[0].column_letter].width = min(60, max(14, width))
                for cell in column[1:]:
                    if isinstance(cell.value, float):
                        cell.number_format = '#,##0.00;[Red](#,##0.00)'
            if len(table.columns) == 1:
                ws.column_dimensions['A'].width = 110
                for row in ws.iter_rows(min_row=2):
                    row[0].alignment = Alignment(wrap_text=True, vertical='top')
                    ws.row_dimensions[row[0].row].height = 45

    print(f'Exported to {TARGET_XLSX.name}: {", ".join(sheets)}')


query_notes = {}


def send_notes_to_xlsx(query_number, notes, assumptions):
    """Refresh the shared notes sheet after each query, preserving other sections."""
    query_notes[query_number] = {
        'Assumptions & Definitions': assumptions,
        'Notes': notes.iloc[:, 0].tolist(),
    }
    wb = load_workbook(TARGET_XLSX)
    for name in ['Query 6 Notes', 'Query 8 Comments', 'Notes & Assumptions']:
        if name in wb.sheetnames:
            del wb[name]
    ws = wb.create_sheet('Notes & Assumptions')
    ws.column_dimensions['A'].width = 110
    ws.append(['Notes & Assumptions'])
    ws['A1'].font = Font(bold=True, size=14)
    ws.row_dimensions[1].height = 24
    ws.freeze_panes = 'A2'
    for number, sections in sorted(query_notes.items()):
        ws.append([None])
        ws.append([f'Query {number}'])
        ws.cell(ws.max_row, 1).font = Font(bold=True, size=12)
        ws.row_dimensions[ws.max_row].height = 22
        for heading, paragraphs in sections.items():
            ws.append([heading])
            ws.cell(ws.max_row, 1).font = Font(bold=True, size=11)
            ws.row_dimensions[ws.max_row].height = 22
            for paragraph in paragraphs:
                ws.append([paragraph])
                ws.cell(ws.max_row, 1).alignment = Alignment(wrap_text=True, vertical='top')
                ws.row_dimensions[ws.max_row].height = max(45, ((len(paragraph) + 89) // 90) * 15 + 15)
            ws.append([None])
    wb.save(TARGET_XLSX)
    print(f'Exported to {TARGET_XLSX.name}: Notes & Assumptions')


def send_fig_to_xlsx(fig, sheet_name, anchor='A1'):
    """Add the PNG to the query's existing sheet, keeping its table."""
    if not TARGET_XLSX.exists():
        raise FileNotFoundError('Run the setup block before exporting a chart.')
    buf = BytesIO()
    fig.savefig(buf, format='png', dpi=150, bbox_inches='tight')
    buf.seek(0)
    wb = load_workbook(TARGET_XLSX)
    ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.create_sheet(sheet_name)
    img = XLImage(buf)
    img.height = img.height * 1000 / img.width
    img.width = 1000
    ws.add_image(img, anchor)
    wb.save(TARGET_XLSX)
    print(f'Exported to {TARGET_XLSX.name}: {sheet_name}')


# Import the source data and check the columns used in the analysis.
df = pd.read_excel(SOURCE_XLSX, sheet_name='Final-Version')
required = [
    'Trade Year', 'Trade Month', CUSTOMER, SECTOR, 'Trading Business Area',
    'Parent Product', 'Currency Code', REVENUE, VOLUME,
]
if df[required].isna().any().any():
    raise ValueError('There are missing values in columns required for the analysis.')
if not np.isfinite(df[[REVENUE, VOLUME]].to_numpy(dtype=float)).all():
    raise ValueError('Revenue and volume must contain finite numbers.')
if (df[VOLUME] < 0).any():
    raise ValueError('Negative volumes require a separate gross/net volume decision.')
if set(df['Trade Year']) != set(YEARS) or not df['Trade Month'].isin(MONTHS).all():
    raise ValueError('Unexpected year or month labels.')

# Use the monthly USD/ZAR table to convert Volume USD to Volume ZAR.
monthly_commercial_market_avg_exchange_rates = pd.read_csv(RATES_CSV)
rates = monthly_commercial_market_avg_exchange_rates
if rates[['Year', 'Month']].duplicated().any():
    raise ValueError('The exchange-rate table has duplicate month and year keys.')
if rates[['Year', 'Month', 'USD/ZAR']].isna().any().any():
    raise ValueError('The exchange-rate table is incomplete.')
if not (np.isfinite(rates['USD/ZAR']) & (rates['USD/ZAR'] > 0)).all():
    raise ValueError('Exchange rates must be positive finite values.')

valued = df.merge(
    rates,
    left_on=['Trade Year', 'Trade Month'],
    right_on=['Year', 'Month'],
    how='left',
    validate='many_to_one',
)
if valued['USD/ZAR'].isna().any():
    raise ValueError('Some source rows did not match an exchange rate.')

# USD/ZAR is rands per dollar, so Volume ZAR = Volume USD * USD/ZAR.
valued['Volume ZAR'] = valued[VOLUME] * valued['USD/ZAR']
check_close(valued[REVENUE].sum(), df[REVENUE].sum(), 'rate join revenue')
annual_totals = df.groupby('Trade Year')[[REVENUE, VOLUME]].sum()

# Start with a clean target workbook every time the setup block is run.
methodology = pd.DataFrame({'Note': [
    'Sources: GM-Case-study-Data-2026.xlsx; South African Revenue Service (SARS), Average Exchange Rates, Table B: '
    'https://www.sars.gov.za/wp-content/uploads/Legal/Rates/Legal-Pub-AER-03-Average-Exchange-Rates-Table-B.pdf '
    '(monthly USD/ZAR data stored in Source-Material/exchange-rates-zar-2020-2022.csv).',
    'Volume ZAR = Volume USD multiplied by the monthly USD/ZAR rate for the trade month.',
    'Margin = summed Client Value ZAR / summed Volume ZAR * 10,000. This is to arrive at basis points (bps) for comparison.',
    'Query 2 treats a missing client-year as zero recorded revenue and ranks changes by ZAR.',
    'Query 4 includes zero-revenue volume and compares clients with defined margins in both years.',
    'Query 6 calculates sector margins from positive-volume records only. Excluded zero-volume revenue remains included in revenue-contribution analyses.',
    'Query 7 explains the exporter-sector method without calculations.',
    'Query 8 includes CNY in either currency field and counts each source row once.',
]})
TARGET_XLSX.parent.mkdir(parents=True, exist_ok=True)
with pd.ExcelWriter(TARGET_XLSX, engine='openpyxl', mode='w') as writer:
    methodology.to_excel(writer, sheet_name='Methodology', index=False)
print(f'Created clean output workbook: {TARGET_XLSX.name}')


# %% QUERY 1: Top 5 highest and lowest revenue contribution clients in 2022
# Group 2022 by Customer Name and sum Client Value ZAR.
# The lowest clients exclude net totals of 0 after positive and negative rows are added together.
df_2022 = df[df['Trade Year'] == 2022]
clients_2022 = df_2022.groupby(CUSTOMER)[REVENUE].sum()
client_dims_2022 = client_dimensions(df_2022)


def clients_2022_table(values):
    return values.rename(REVENUE).to_frame().join(client_dims_2022).reset_index()


top_5_clients_2022 = clients_2022_table(
    clients_2022.sort_values(ascending=False).head(5)
)
bottom_5_clients_2022 = clients_2022_table(
    clients_2022[clients_2022 != 0].sort_values().head(5)
)
top_5_clients_2022 = number_rows(top_5_clients_2022)
bottom_5_clients_2022 = number_rows(bottom_5_clients_2022)
print_table('Top 5 clients in 2022', top_5_clients_2022)
print_table('Bottom 5 nonzero clients in 2022', bottom_5_clients_2022)
send_to_xlsx({
    'Query 1 T5 Revenue Clients 2022': top_5_clients_2022[['Rank', CUSTOMER, REVENUE]],
    'Query 1 B5 Revenue Clients 2022': bottom_5_clients_2022[['Rank', CUSTOMER, REVENUE]],
})


# %% QUERY 2: Top 5 revenue gainers and droppers from 2021 to 2022
# Rank on the ZAR change. Percentage change is included as context.
# Percentage change is blank where 2021 revenue is 0 or negative.
products_yoy = df[df['Trade Year'].isin([2021, 2022])]
client_revenue_yoy = (
    products_yoy.groupby([CUSTOMER, 'Trade Year'])[REVENUE]
    .sum().unstack().reindex(columns=[2021, 2022])
)
client_revenue_yoy['Coverage'] = np.select(
    [client_revenue_yoy[2021].isna(), client_revenue_yoy[2022].isna()],
    ['Only recorded in 2022', 'Only recorded in 2021'],
    default='Recorded in both years',
)
client_revenue_yoy[[2021, 2022]] = client_revenue_yoy[[2021, 2022]].fillna(0)
client_revenue_yoy['Revenue change ZAR'] = client_revenue_yoy[2022] - client_revenue_yoy[2021]
client_revenue_yoy['Revenue change (%)'] = (
    client_revenue_yoy['Revenue change ZAR']
    / client_revenue_yoy[2021].where(client_revenue_yoy[2021] > 0) * 100
)
client_revenue_yoy = client_revenue_yoy.rename(columns={
    2021: 'Client Value ZAR 2021', 2022: 'Client Value ZAR 2022',
})
client_revenue_yoy = client_revenue_yoy.join(client_dimensions(products_yoy)).reset_index()

delta_gainers = (
    client_revenue_yoy[client_revenue_yoy['Revenue change ZAR'] > 0]
    .sort_values(['Revenue change ZAR', CUSTOMER], ascending=[False, True]).head(5)
)
delta_droppers = (
    client_revenue_yoy[client_revenue_yoy['Revenue change ZAR'] < 0]
    .sort_values(['Revenue change ZAR', CUSTOMER]).head(5)
)
check_close(
    client_revenue_yoy['Revenue change ZAR'].sum(),
    annual_totals.loc[2022, REVENUE] - annual_totals.loc[2021, REVENUE],
    'client changes to total revenue change',
)
delta_gainers = number_rows(delta_gainers)
delta_droppers = number_rows(delta_droppers)
print_table('Top 5 revenue gainers by ZAR change', delta_gainers)
print_table('Top 5 revenue droppers by ZAR change', delta_droppers)
revenue_change_cols = ['Rank', CUSTOMER, 'Client Value ZAR 2021', 'Client Value ZAR 2022',
                       'Revenue change ZAR', 'Revenue change (%)']
send_to_xlsx({
    'Query 2 T5 Gainers': delta_gainers[revenue_change_cols],
    'Query 2 T5 Droppers': delta_droppers[revenue_change_cols],
})


# %% QUERY 3: Monthly run rate analysis from 2020 to 2022
# Sum Client Value ZAR by Trade Month and Trade Year and use calendar month order.
monthly_value = (
    df.groupby(['Trade Month', 'Trade Year'])[REVENUE].sum()
    .unstack('Trade Year').reindex(index=MONTHS, columns=YEARS)
)
if monthly_value.isna().any().any():
    raise ValueError('At least one month-year combination is missing.')
check_close(
    monthly_value.sum().to_numpy(),
    annual_totals.loc[YEARS, REVENUE].to_numpy(),
    'monthly revenue to annual totals',
)

x = np.arange(len(MONTHS))
bar_width = 0.18
offsets = (np.arange(len(YEARS)) - 1) * bar_width
fig, ax = plt.subplots()
for year, offset in zip(YEARS, offsets):
    ax.bar(
        x + offset, monthly_value[year], width=bar_width,
        color=CHART_COLORS['year'][year], label=str(year), linewidth=0, zorder=3,
    )
ax.set_xticks(x)
ax.set_xticklabels([month[:3] for month in MONTHS])
ax.set_xlim(-0.6, len(MONTHS) - 0.4)
ax.set_ylabel('Client Value (ZAR million)')
ax.set_title('Monthly run rate analysis YoY for 2020 to 2022')
ax.yaxis.set_major_formatter(FuncFormatter(lambda value, _: f'{value / 1e6:.0f}'))
ax.legend(ncol=3, loc='upper right', bbox_to_anchor=(1, 1.14), frameon=False)
fig.tight_layout()

monthly_output = monthly_value.rename_axis('Trade Month').reset_index()
print_table('Monthly run rate analysis', monthly_output)
send_to_xlsx({'Query 3 Monthly Revenue': monthly_output})
send_fig_to_xlsx(fig, 'Query 3 Monthly Revenue', anchor='F2')
plt.show(block=False)


# %% QUERY 4: Top 10 FX client margin gainers and droppers from 2021 to 2022
# Calculate margin from summed revenue and summed converted volume.
# Zero-revenue volume remains in the denominator.
fx = valued[valued['Trading Business Area'] == 'FOREIGN EXCHANGE'].copy()
fx_summary = (
    fx.groupby(['Trade Year', CUSTOMER])[[REVENUE, 'Volume ZAR']]
    .sum().reset_index()
)
fx_summary['Margin (bps)'] = margin_bps(fx_summary[REVENUE], fx_summary['Volume ZAR'])
fx_customers_2021 = fx_summary[fx_summary['Trade Year'] == 2021].drop(columns='Trade Year')
fx_customers_2022 = fx_summary[fx_summary['Trade Year'] == 2022].drop(columns='Trade Year')
fx_margin_coverage = fx_customers_2021.merge(
    fx_customers_2022, on=CUSTOMER, how='outer',
    suffixes=('_2021', '_2022'), validate='one_to_one', indicator=True,
)
fx_margin_coverage['Comparable'] = (
    fx_margin_coverage[['Margin (bps)_2021', 'Margin (bps)_2022']]
    .notna().all(axis=1)
)
fx_margin_yoy = (
    fx_margin_coverage[fx_margin_coverage['Comparable']].copy()
    .rename(columns={
        'Margin (bps)_2021': 'Margin 2021 (bps)',
        'Margin (bps)_2022': 'Margin 2022 (bps)',
    })
)
fx_margin_yoy['Margin delta (bps)'] = (
    fx_margin_yoy['Margin 2022 (bps)'] - fx_margin_yoy['Margin 2021 (bps)']
)
fx_sector = (
    fx[fx['Trade Year'].isin([2021, 2022])]
    .groupby(CUSTOMER)[SECTOR].agg(join_unique)
)
fx_margin_yoy = fx_margin_yoy.join(fx_sector.rename('Sector'), on=CUSTOMER)
margin_yoy_cols = [
    CUSTOMER, 'Margin 2021 (bps)', 'Margin 2022 (bps)', 'Margin delta (bps)',
]
fx_margin_gainers = (
    fx_margin_yoy.loc[fx_margin_yoy['Margin delta (bps)'] > 0, margin_yoy_cols]
    .sort_values(['Margin delta (bps)', CUSTOMER], ascending=[False, True]).head(10)
)
fx_margin_droppers = (
    fx_margin_yoy.loc[fx_margin_yoy['Margin delta (bps)'] < 0, margin_yoy_cols]
    .sort_values(['Margin delta (bps)', CUSTOMER]).head(10)
)
fx_margin_gainers = number_rows(fx_margin_gainers)
fx_margin_droppers = number_rows(fx_margin_droppers)
print_table('Top 10 FX margin gainers', fx_margin_gainers)
print_table('Top 10 FX margin droppers', fx_margin_droppers)
send_to_xlsx({
    'Query 4 T10 FX Gainers': fx_margin_gainers,
    'Query 4 B10 FX Droppers': fx_margin_droppers,
})


# %% QUERY 5: Revenue contribution change by Parent Product from 2021 to 2022
# Calculate each product's share against total revenue for the same year.
# The difference between the two shares is a percentage-point change.
products_yoy = df[df['Trade Year'].isin([2021, 2022])]
parent_product_revenue = (
    products_yoy.groupby(['Parent Product', 'Trade Year'])[REVENUE]
    .sum().unstack('Trade Year').reindex(columns=[2021, 2022]).fillna(0)
)
for year in [2021, 2022]:
    total = annual_totals.loc[year, REVENUE]
    if total == 0:
        raise ValueError(f'Cannot calculate product revenue shares for {year}.')
    parent_product_revenue[f'Revenue Contribution {year} (%)'] = (
        parent_product_revenue[year] / total * 100
    )
    check_close(parent_product_revenue[year].sum(), total, f'{year} product revenue')
    check_close(
        parent_product_revenue[f'Revenue Contribution {year} (%)'].sum(),
        100, f'{year} product shares',
    )
parent_product_revenue['Revenue Contribution delta (pp)'] = (
    parent_product_revenue['Revenue Contribution 2022 (%)']
    - parent_product_revenue['Revenue Contribution 2021 (%)']
)
parent_product_revenue = (
    parent_product_revenue
    .rename(columns={2021: 'Client Value ZAR 2021', 2022: 'Client Value ZAR 2022'})
    .sort_values('Revenue Contribution delta (pp)', ascending=False).reset_index()
)
print_table('Parent product revenue contribution', parent_product_revenue)
send_to_xlsx({'Query 5 Product Revenue': parent_product_revenue[[
    'Parent Product', 'Revenue Contribution 2021 (%)',
    'Revenue Contribution 2022 (%)', 'Revenue Contribution delta (pp)',
]]})


# %% QUERY 6: Relationship between margin and sector in 2022
# Use positive-volume records for both revenue and volume in all sector margins.
sector_margin_source = valued[valued['Trade Year'] == 2022].copy()


def sector_totals(source):
    totals = source.groupby(SECTOR).agg(
        **{
            REVENUE: (REVENUE, 'sum'),
            'Volume ZAR': ('Volume ZAR', 'sum'),
            'Records': (CUSTOMER, 'size'),
            'Clients': (CUSTOMER, 'nunique'),
        }
    )
    totals.index.name = 'Sector'
    totals['Margin (bps)'] = margin_bps(totals[REVENUE], totals['Volume ZAR'])
    return totals


positive_volume = sector_margin_source[sector_margin_source[VOLUME] > 0]
zero_volume = sector_margin_source[sector_margin_source[VOLUME] == 0]
sector_margin_2022 = sector_totals(positive_volume).rename(columns={
    'Margin (bps)': 'Positive-volume margin (bps)'
})
fx_only_source = positive_volume[
    positive_volume['Trading Business Area'] == 'FOREIGN EXCHANGE'
]
sector_margin_2022['FX only margin (bps)'] = sector_totals(fx_only_source)['Margin (bps)']
sector_margin_2022['Zero-volume revenue ZAR'] = (
    zero_volume.groupby(SECTOR)[REVENUE].sum()
    .reindex(sector_margin_2022.index, fill_value=0)
)
money_market_volume = (
    sector_margin_source[
        sector_margin_source['Trading Business Area'] == 'MONEY MARKETS'
    ].groupby(SECTOR)['Volume ZAR'].sum()
)
sector_margin_2022['Money Markets volume share (%)'] = (
    money_market_volume.reindex(sector_margin_2022.index, fill_value=0)
    / sector_margin_2022['Volume ZAR'].where(sector_margin_2022['Volume ZAR'] > 0) * 100
)
check_close(
    sector_margin_2022[REVENUE].sum() + zero_volume[REVENUE].sum(),
    annual_totals.loc[2022, REVENUE],
    'positive-volume sector revenue plus excluded zero-volume revenue',
)
sector_margin_2022 = (
    sector_margin_2022.sort_values('Positive-volume margin (bps)', ascending=False)
    .reset_index()
)
zero_volume_summary = (
    zero_volume.groupby([SECTOR, 'Trading Business Area'])
    .agg(Records=(REVENUE, 'size'), **{REVENUE: (REVENUE, 'sum')})
    .reset_index()
)
sector_product = (
    positive_volume.groupby([SECTOR, 'Parent Product'])[[REVENUE, 'Volume ZAR']]
    .sum().reset_index()
)
sector_product['Margin (bps)'] = margin_bps(
    sector_product[REVENUE], sector_product['Volume ZAR']
)
sector_product['Within-sector volume share (%)'] = (
    sector_product['Volume ZAR']
    / sector_product.groupby(SECTOR)['Volume ZAR'].transform('sum').replace(0, np.nan) * 100
)

retail = positive_volume[positive_volume[SECTOR] == 'UNKNOWN SECTOR']
highest = sector_margin_2022.dropna(subset=['Positive-volume margin (bps)']).iloc[0]
lowest = sector_margin_2022.dropna(subset=['Positive-volume margin (bps)']).iloc[-1]
fx_lowest = (
    sector_margin_2022.dropna(subset=['FX only margin (bps)'])
    .sort_values('FX only margin (bps)').iloc[0]
)
sector_margin_assumptions = [
    'Scope: 2022. UNKNOWN SECTOR is treated as retail clients.',
    'Margin (bps): sector revenue across all included products (bonds, overnight deposits, term deposits, FX forwards, FX options, FX spot, FX swaps and vanilla equity options) divided by sector Volume ZAR, multiplied by 10,000. Both totals use only records with positive recorded USD volume in 2022.',
    'FX only margin (bps): the same calculation restricted to Trading Business Area = FOREIGN EXCHANGE (FX spot, forwards, swaps and options), again using only positive-volume records.',
    'Why show FX only: sectors have different mixes of FX, deposits and other products, whose revenue-to-volume ratios are not directly comparable. Restricting the comparison to FX helps show whether a low overall margin reflects non-FX product mix. For example, financial institutions have 1.19 bps across all included products versus 10.75 bps for FX only; 95.8% of their recorded volume is money markets. FX-only margins can still differ because of FX product mix, deal sizes and client composition, so they do not isolate pricing alone.',
    'Volume ZAR = Volume USD multiplied by the monthly USD/ZAR exchange rate. One basis point (bp) equals 0.01 percentage points.',
    'Client Value ZAR, Volume ZAR, Records and Clients in the Query 6 table cover the included positive-volume records across all products. Records counts source rows; Clients counts distinct customer names.',
    'Some records contain revenue but zero recorded USD volume, which also produces zero '
    'converted ZAR volume. These occur in money-market products and, for mining, '
    'predominantly in products classified as Commodities.',
    'The dataset does not explain whether this reflects missing volume, a reporting convention, '
    'or revenue booked separately from the underlying transaction. I therefore excluded those '
    'records from sector-margin calculations because their revenue-to-volume ratio is undefined, '
    'while retaining their revenue in the revenue-contribution analysis.',
    'Monthly USD/ZAR conversion (based on data from the South African Revenue Service…table B) remains an approximation.',
]

sector_margin_2022_notes = pd.DataFrame({'Note': [
    f"Highest sector margin: {highest['Sector']}, "
    f"{highest['Positive-volume margin (bps)']:.2f} bps. "
    f"Lowest: {lowest['Sector']}, {lowest['Positive-volume margin (bps)']:.2f} bps.",
    'Sector margins reflect product mix as well as pricing; they are not a controlled '
    'comparison of the prices charged to different sectors.',
    f"{lowest['Sector']} has {lowest['Money Markets volume share (%)']:.1f}% money-market "
    f"volume. The lowest FX-only sector is {fx_lowest['Sector']} at "
    f"{fx_lowest['FX only margin (bps)']:.2f} bps.",
    'Higher retail margins could reflect smaller deal sizes and less bargaining power '
    '(Retail clients are price takers). Lower financial-institution margins reflect informed flows, '
    'and their larger volumes mean stronger price negotiation. These are both possible explanations, '
    'not conclusions established by this sample.',
    f"UNKNOWN SECTOR is retail. It contains {len(retail)} record(s), "
    f"{retail[CUSTOMER].nunique()} client(s) and USD {retail[VOLUME].sum():,.2f} volume. "
    'This is too small a sample to represent retail clients generally.',
]})

sector_margin_output = sector_margin_2022[[
    'Sector', 'Client Value ZAR', 'Volume ZAR', 'Records', 'Clients',
    'Positive-volume margin (bps)', 'FX only margin (bps)',
]].rename(columns={'Positive-volume margin (bps)': 'Margin (bps)'}).copy()
sector_margin_output['Sector'] = sector_margin_output['Sector'].replace({
    'UNKNOWN SECTOR': 'RETAIL (UNKNOWN SECTOR)',
})
print_table('2022 sector margins', sector_margin_output)
print_table('Sector notes', sector_margin_2022_notes)
send_to_xlsx({
    'Query 6 Sector Margins': sector_margin_output,
})
send_notes_to_xlsx(6, sector_margin_2022_notes, sector_margin_assumptions)


# %% QUERY 7: Method for identifying the largest exporter sector by volume
# Explain the method without calculating a sector ranking.
exporter_method = pd.DataFrame({'Step': [
    'Assumption: exporter-shaped behavior means a domestic client converting inbound '
    'foreign currency into ZAR. Treat Buy/Sell Flag as the client buying or selling '
    'the currency in Currency Code.',
    '1. Select every row where a foreign currency trades against ZAR, with ZAR in '
    'either Currency Code or Counter Currency Code and a foreign currency in the other field.',
    '2. Sign Volume USD: positive when the client sells foreign currency for ZAR; '
    'negative when the client buys foreign currency with ZAR. If Currency Code is '
    'foreign and Counter Currency Code is ZAR, S is positive and B is negative. '
    'If Currency Code is ZAR and Counter Currency Code is foreign, B is positive and S is negative.',
    '3. Sum signed Volume USD per Customer Name across the whole period (2020 to 2022). '
    'A positive net identifies a net converter of foreign currency into ZAR '
    '(exporter-shaped behavior); a negative net indicates importer-shaped behavior. '
    'A zero net is neither.',
    '4. Keep only positive-net customers, tag each with Immediate Parent Client Sector, '
    'sum their positive net USD-equivalent volumes by sector, and rank sectors from '
    'largest to smallest. Rank by net conversion volume, not customer count or gross trade volume.',
]})
print_table('Method for identifying exporter sectors', exporter_method)
send_to_xlsx({'Query 7 Exporter Method': exporter_method})


# %% QUERY 8: CNY trade volume from 2020 to 2022
# Include records where CNY appears in Currency Code or Counter Currency Code.
# The figures are USD-equivalent trade volumes, not amounts denominated in CNY.
cny = df[
    (df['Currency Code'] == 'CNY') | (df['Counter Currency Code'] == 'CNY')
]
cny_products = sorted(cny['Parent Product'].unique())
cny_rows = []
for year in YEARS:
    year_cny = cny[cny['Trade Year'] == year]
    total_volume = annual_totals.loc[year, VOLUME]
    volume_usd = year_cny[VOLUME].sum()
    row = {
        'Year': year,
        'Currency involvement': 'CNY in either field',
        VOLUME: volume_usd,
        'Share of Order Book USD volume (%)': (
            volume_usd / total_volume * 100 if total_volume > 0 else np.nan
        ),
    }
    for product in cny_products:
        row[product] = year_cny.loc[
            year_cny['Parent Product'] == product, VOLUME
        ].sum()
    check_close(
        sum(row[product] for product in cny_products),
        volume_usd,
        f'{year} CNY product totals',
    )
    cny_rows.append(row)

join_cny_volume_2020_2021_2022 = pd.DataFrame(cny_rows)
previous_volume = join_cny_volume_2020_2021_2022[VOLUME].shift()
join_cny_volume_2020_2021_2022['Volume YoY change (%)'] = (
    (join_cny_volume_2020_2021_2022[VOLUME] - previous_volume)
    / previous_volume.where(previous_volume > 0) * 100
)
cny_volume_2022 = (
    join_cny_volume_2020_2021_2022.set_index('Year').loc[2022, VOLUME]
)
cny_output = join_cny_volume_2020_2021_2022[[
    'Year', 'Currency involvement', *cny_products, VOLUME,
    'Share of Order Book USD volume (%)', 'Volume YoY change (%)',
]]
print_table('CNY trade volume by year', cny_output)
print(f'\n2022 CNY trade volume: USD {cny_volume_2022:,.2f}')

# Less than 1% of the order book traded CNY in 2020; this increased to 13.07% in 2022.
# More diversification of products could suggest increased trading with China.
# Possible explanation: China's lockdown and reopening timing, including the idea
# that it entered and exited lockdown earlier, may have affected demand and exports
# to China. The timing and any export-volume link need external evidence.
# Possible explanation: growing BRICS trade ties may have supported CNY activity.
# The dataset doesn't establish either explanation or identify export destinations.
# CNY-related 2022 trade volume = USD 457.6M, including CNY in either currency field.
cny_by_year = join_cny_volume_2020_2021_2022.set_index('Year')
cny_assumptions = [
    'Scope: 2020 to 2022. We included a source row when CNY appears in either Currency Code or Counter Currency Code, and counted each row once.',
    'CNY trade volume is the sum of recorded Volume USD for those rows: USD-equivalent volume, not an amount denominated in CNY.',
    'Order-book share = CNY-related USD volume divided by total recorded portfolio USD volume in the same year, multiplied by 100.',
    'Volume YoY change = (current-year CNY volume / prior-year CNY volume - 1) multiplied by 100.',
    'CNY currency involvement does not identify an explicit export destination or prove underlying trade with China. COVID Lockdown and BRICS explanations are hypotheses requiring independent verification.',
]
cny_comments = pd.DataFrame({'Comment': [
    f"Less than 1% of order-book USD volume involved CNY in 2020 "
    f"({cny_by_year.loc[2020, 'Share of Order Book USD volume (%)']:.2f}%), "
    f"increasing to {cny_by_year.loc[2022, 'Share of Order Book USD volume (%)']:.2f}% in 2022. "
    'Here, order-book share means share of total recorded USD trade volume.',
    'We see diversification of CNY products with FX swaps being included in 2022. '
    'This supports the thesis of increased trading with China. '
    'It could also reflect changes in hedging or financing.',
    "Possible explanation 1: China's lockdown and reopening timing, including the suggestion "
    'that it entered and exited lockdown earlier, may have contributed to a demand surge '
    'and higher exports to China. The timing and export-volume link require further inspection.',
    'Possible explanation 2: BRICS trade ties may have been gaining momentum during this '
    'period and supporting CNY activity.',
    f'CNY-related 2022 trade volume = USD {cny_volume_2022 / 1e6:,.1f} million '
    f'(USD {cny_volume_2022:,.2f}), including rows where CNY appears in either '
    'Currency Code or Counter Currency Code, with each row counted once.',
]})
print_table('Query 8 comments', cny_comments)
send_to_xlsx({
    'Query 8 CNY Volume': cny_output,
})
send_notes_to_xlsx(8, cny_comments, cny_assumptions)
